"""Convert Markdown notes/translations to Word (.docx) for machines without MD readers.

Formulas are rendered in a human-readable Word style (native sub/superscripts,
Cambria Math, optional OMML fractions) — not raw LaTeX or underscore text.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from src.formula_humanize import (
    add_formula_runs,
    humanize_inline_text,
    looks_like_formula,
)

HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
UL_RE = re.compile(r"^(\s*)([-*+])\s+(.*)$")
OL_RE = re.compile(r"^(\s*)(\d+)[.)]\s+(.*)$")
HR_RE = re.compile(r"^(\*{3,}|-{3,}|_{3,})\s*$")
TABLE_ROW_RE = re.compile(r"^\|(.+)\|\s*$")
TABLE_SEP_RE = re.compile(r"^\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$")
IMAGE_RE = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")
LINK_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)")
BOLD_ITALIC_RE = re.compile(
    r"(\*\*\*[^*]+?\*\*\*|\*\*[^*]+?\*\*|\*[^*]+?\*|___[^_]+?___|__[^_]+?__|`[^`]+`)"
)
# Embedded formula-ish tokens inside prose (avoid matching markdown emphasis _)
INLINE_FORMULA_RE = re.compile(
    r"(?<![A-Za-z0-9])("
    r"[A-Za-zΑ-Ωα-ωεμσΔ∂πΠ][A-Za-z0-9]*"
    r"_(?:\{[^}]+\}|[A-Za-z0-9]+(?:/[A-Za-z0-9]+)?(?:,[A-Za-z0-9]+)?)"
    r"(?:\^(?:\{[^}]+\}|[A-Za-z0-9+-]+))?"
    r"|"
    r"[A-Za-zΑ-Ωα-ωεμσΔ∂πΠ][A-Za-z0-9]*\^(?:\{[^}]+\}|[A-Za-z0-9+-]+)"
    r"|"
    r"\d+\^(?:\{[^}]+\}|-?\d+)"
    r"|"
    r"10\^\{?-?\d+\}?"
    r")(?![A-Za-z0-9])"
)


def _set_run_font(
    run,
    *,
    east_asia: str = "宋体",
    ascii_font: str = "Calibri",
    size_pt: float = 11,
) -> None:
    run.font.size = Pt(size_pt)
    run.font.name = ascii_font
    r = run._element
    rpr = r.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    rfonts.set(qn("w:ascii"), ascii_font)
    rfonts.set(qn("w:hAnsi"), ascii_font)
    rfonts.set(qn("w:eastAsia"), east_asia)


def _add_plain_segment(paragraph, text: str, *, base_size: float = 11) -> None:
    """Markdown bold/italic/code for ordinary prose (no underscore math)."""
    if not text:
        return
    text = humanize_inline_text(text)
    pos = 0
    for match in BOLD_ITALIC_RE.finditer(text):
        if match.start() > pos:
            chunk = text[pos : match.start()]
            _add_prose_with_inline_formulas(paragraph, chunk, base_size=base_size)
        token = match.group(0)
        if token.startswith("***") and token.endswith("***"):
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
            _set_run_font(run, size_pt=base_size)
        elif token.startswith("**") and token.endswith("**"):
            inner = token[2:-2]
            if looks_like_formula(inner):
                # bold wrapper around a formula — still use math formatting
                add_formula_runs(paragraph, inner, size_pt=base_size)
                for run in paragraph.runs[-3:]:
                    run.bold = True
            else:
                run = paragraph.add_run(inner)
                run.bold = True
                _set_run_font(run, size_pt=base_size)
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
            _set_run_font(run, size_pt=base_size)
        elif token.startswith("___") and token.endswith("___"):
            run = paragraph.add_run(token[3:-3])
            run.bold = True
            run.italic = True
            _set_run_font(run, size_pt=base_size)
        elif token.startswith("__") and token.endswith("__"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            _set_run_font(run, size_pt=base_size)
        elif token.startswith("`") and token.endswith("`"):
            inner = token[1:-1]
            if looks_like_formula(inner) or "_" in inner or "^" in inner:
                add_formula_runs(paragraph, inner, size_pt=base_size)
            else:
                run = paragraph.add_run(inner)
                _set_run_font(
                    run, ascii_font="Consolas", east_asia="宋体", size_pt=base_size
                )
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        else:
            run = paragraph.add_run(token)
            _set_run_font(run, size_pt=base_size)
        pos = match.end()
    if pos < len(text):
        _add_prose_with_inline_formulas(paragraph, text[pos:], base_size=base_size)


def _add_prose_with_inline_formulas(
    paragraph, text: str, *, base_size: float = 11
) -> None:
    """Inside prose, lift g_abs / A_L / 10^{-3} into proper Word scripts."""
    if not text:
        return
    pos = 0
    for m in INLINE_FORMULA_RE.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            _set_run_font(run, size_pt=base_size)
        add_formula_runs(paragraph, m.group(1), size_pt=base_size)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        _set_run_font(run, size_pt=base_size)


def _add_inline_runs(paragraph, text: str, *, base_size: float = 11) -> None:
    """Route whole-line formulas to math renderer; otherwise prose + inline math."""
    if not text:
        return
    cleaned = humanize_inline_text(text)
    # Whole paragraph / cell is an equation
    if looks_like_formula(cleaned) and not re.search(r"[\u4e00-\u9fff]{12,}", cleaned):
        add_formula_runs(paragraph, cleaned, size_pt=base_size)
        return
    _add_plain_segment(paragraph, text, base_size=base_size)


def _normalize_inline(text: str) -> str:
    """Replace image/link markdown with readable plain text before inline parse."""
    text = IMAGE_RE.sub(lambda m: f"[图片: {m.group(1) or m.group(2)}]", text)
    text = LINK_RE.sub(lambda m: f"{m.group(1)} ({m.group(2)})", text)
    return text


def _split_table_cells(line: str) -> list[str]:
    inner = line.strip()
    if inner.startswith("|"):
        inner = inner[1:]
    if inner.endswith("|"):
        inner = inner[:-1]
    return [c.strip() for c in inner.split("|")]


def _add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell_text = row[j] if j < len(row) else ""
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            _add_inline_runs(p, _normalize_inline(cell_text), base_size=10)
            if i == 0:
                for run in p.runs:
                    run.bold = True


def _add_formula_block(doc: Document, text: str) -> None:
    """Centered display formula, slightly larger."""
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    add_formula_runs(p, text, size_pt=12)


def markdown_to_docx(md_text: str, output_path: Path, *, title: str | None = None) -> Path:
    """Convert Markdown string to a .docx file."""
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    pf = style.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    pf.space_after = Pt(6)

    if title:
        h = doc.add_heading(title, level=0)
        for run in h.runs:
            _set_run_font(run, east_asia="黑体", size_pt=18)
            run.bold = True

    lines = md_text.replace("\r\n", "\n").replace("\r", "\n").split("\n")
    i = 0
    in_code = False
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]

        # Fenced code block
        if line.strip().startswith("```"):
            if not in_code:
                in_code = True
                code_lines = []
            else:
                in_code = False
                block = "\n".join(code_lines)
                # Formula / symbol blocks → human math; otherwise monospace
                if looks_like_formula(block) or (
                    len(code_lines) <= 6
                    and any(c in block for c in "=_^×·αβγδε")
                    and not any(
                        kw in block
                        for kw in ("def ", "import ", "return ", "class ", "{", "}")
                    )
                ):
                    for fl in code_lines:
                        if fl.strip():
                            _add_formula_block(doc, fl.strip())
                else:
                    p = doc.add_paragraph()
                    run = p.add_run(block)
                    _set_run_font(
                        run, ascii_font="Consolas", east_asia="宋体", size_pt=9
                    )
                    p.paragraph_format.left_indent = Inches(0.2)
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if HR_RE.match(line.strip()):
            p = doc.add_paragraph("─" * 30)
            for run in p.runs:
                _set_run_font(run, size_pt=10)
                run.font.color.rgb = RGBColor(0x99, 0x99, 0x99)
            i += 1
            continue

        hm = HEADING_RE.match(line)
        if hm:
            level = min(len(hm.group(1)), 4)
            text = _normalize_inline(hm.group(2).strip())
            text = re.sub(r"\s*\{#.*\}\s*$", "", text)
            h = doc.add_heading(level=level)
            _add_inline_runs(h, text, base_size=16 - level)
            for run in h.runs:
                run.bold = True
            i += 1
            continue

        if TABLE_ROW_RE.match(line) and i + 1 < len(lines) and TABLE_SEP_RE.match(
            lines[i + 1].strip()
        ):
            rows: list[list[str]] = [_split_table_cells(line)]
            i += 2
            while i < len(lines) and TABLE_ROW_RE.match(lines[i]):
                rows.append(_split_table_cells(lines[i]))
                i += 1
            _add_table(doc, rows)
            continue

        if line.lstrip().startswith(">"):
            quote_parts: list[str] = []
            while i < len(lines) and lines[i].lstrip().startswith(">"):
                quote_parts.append(re.sub(r"^>\s?", "", lines[i].lstrip()))
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.25)
            _add_inline_runs(p, _normalize_inline(" ".join(quote_parts)))
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            continue

        um = UL_RE.match(line)
        if um:
            text = _normalize_inline(um.group(3))
            p = doc.add_paragraph(style="List Bullet")
            if looks_like_formula(text) and not re.search(r"[\u4e00-\u9fff]{8,}", text):
                add_formula_runs(p, text, size_pt=11)
            else:
                _add_inline_runs(p, text)
            i += 1
            continue

        om = OL_RE.match(line)
        if om:
            text = _normalize_inline(om.group(3))
            p = doc.add_paragraph(style="List Number")
            if looks_like_formula(text) and not re.search(r"[\u4e00-\u9fff]{8,}", text):
                add_formula_runs(p, text, size_pt=11)
            else:
                _add_inline_runs(p, text)
            i += 1
            continue

        para_text = line
        i += 1
        while i < len(lines):
            nxt = lines[i]
            if (
                not nxt.strip()
                or HEADING_RE.match(nxt)
                or UL_RE.match(nxt)
                or OL_RE.match(nxt)
                or nxt.lstrip().startswith(">")
                or nxt.strip().startswith("```")
                or HR_RE.match(nxt.strip())
                or TABLE_ROW_RE.match(nxt)
            ):
                break
            para_text += " " + nxt.strip()
            i += 1

        cleaned = _normalize_inline(para_text)
        if looks_like_formula(cleaned) and not re.search(
            r"[\u4e00-\u9fff]{12,}", cleaned
        ):
            _add_formula_block(doc, cleaned)
        else:
            p = doc.add_paragraph()
            _add_inline_runs(p, cleaned)

    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def convert_md_file(md_path: Path, docx_path: Path | None = None) -> Path:
    md_path = Path(md_path)
    if not md_path.is_file():
        raise FileNotFoundError(f"找不到 Markdown 文件: {md_path}")
    docx_path = Path(docx_path) if docx_path else md_path.with_suffix(".docx")
    text = md_path.read_text(encoding="utf-8")
    title = md_path.stem
    return markdown_to_docx(text, docx_path, title=title)


NOTE_BASENAMES = ("reading_note", "translation_zh", "translation_si_zh")


def export_notes_dir(notes_dir: Path, *, overwrite: bool = True) -> dict[str, str | None]:
    """
    Export reading_note / translation_zh / translation_si_zh Markdown to .docx.

    Returns map basename -> docx path (or None if md missing).
    """
    notes_dir = Path(notes_dir)
    results: dict[str, str | None] = {}
    for name in NOTE_BASENAMES:
        md = notes_dir / f"{name}.md"
        docx = notes_dir / f"{name}.docx"
        if not md.is_file():
            results[name] = None
            continue
        if docx.is_file() and not overwrite:
            results[name] = str(docx)
            continue
        convert_md_file(md, docx)
        results[name] = str(docx)
    return results
